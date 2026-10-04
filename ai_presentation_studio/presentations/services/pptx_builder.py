from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .images import fetch_pexels_image

W, H = 13.333, 7.5

# Presentation color presets. The website brand colors are independent from these.
THEMES = {
    "red": {"bg": "0B0B0D", "fg": "FFFFFF", "muted": "B9B9BF", "accent": "E10600", "panel": "18181B"},
    "mono": {"bg": "0B0B0D", "fg": "FFFFFF", "muted": "A8A8AE", "accent": "FFFFFF", "panel": "1B1B1F"},
    "ocean": {"bg": "07131F", "fg": "F7FAFF", "muted": "9EB3C9", "accent": "2F80FF", "panel": "102234"},
    "emerald": {"bg": "071411", "fg": "F4FFF9", "muted": "9CC9B7", "accent": "44F3A4", "panel": "10211C"},
    "violet": {"bg": "120B1B", "fg": "F9F5FF", "muted": "C4B5D5", "accent": "B56CFF", "panel": "21142F"},
    "sunset": {"bg": "180C08", "fg": "FFF8F4", "muted": "D8B3A5", "accent": "FF6B35", "panel": "2B1710"},
}


def _normalize_hex(value: str, fallback: str) -> str:
    value = (value or "").strip().lstrip("#").upper()
    if len(value) == 6 and all(ch in "0123456789ABCDEF" for ch in value):
        return value
    return fallback


def _hex_tuple(value: str):
    value = _normalize_hex(value, "000000")
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def _mix_hex(a: str, b: str, ratio: float) -> str:
    ar, ag, ab = _hex_tuple(a)
    br, bg, bb = _hex_tuple(b)
    ratio = max(0.0, min(1.0, ratio))
    vals = (
        round(ar * (1 - ratio) + br * ratio),
        round(ag * (1 - ratio) + bg * ratio),
        round(ab * (1 - ratio) + bb * ratio),
    )
    return "".join(f"{v:02X}" for v in vals)


def _contrast_text(background: str) -> str:
    r, g, b = _hex_tuple(background)
    luminance = (0.299 * r + 0.587 * g + 0.114 * b)
    return "111111" if luminance > 165 else "FFFFFF"


def resolve_palette(preset="red", custom_colors=None):
    """Return a complete palette for preview/PPTX generation."""
    if preset != "custom":
        base = dict(THEMES.get(preset, THEMES["red"]))
    else:
        custom_colors = custom_colors or {}
        bg = _normalize_hex(custom_colors.get("background"), "111111")
        accent = _normalize_hex(custom_colors.get("accent"), "FF0000")
        fg = _normalize_hex(custom_colors.get("text"), "FFFFFF")
        base = {
            "bg": bg,
            "fg": fg,
            "accent": accent,
            "panel": _mix_hex(bg, fg, 0.10),
            "muted": _mix_hex(bg, fg, 0.62),
        }
    base["accent_text"] = _contrast_text(base["accent"])
    return base


def rgb(hexv):
    hexv = hexv.lstrip("#")
    return RGBColor(int(hexv[0:2], 16), int(hexv[2:4], 16), int(hexv[4:6], 16))


def fill_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_text(slide, text, x, y, w, h, size=24, color="FFFFFF", bold=False, align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text or ""
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    return box


def panel(slide, x, y, w, h, color, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_top_tag(slide, t):
    panel(slide, 0.65, 0.48, 1.55, 0.42, t["accent"])
    add_text(slide, "AI PRESENTATION", 0.78, 0.54, 1.3, 0.2, 8, t["accent_text"], True)


def add_bullets(slide, items, x, y, w, h, t, size=21):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items[:6]):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "•  " + str(item)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(t["fg"])
        p.space_after = Pt(14)
    return box


def slide_title(slide, title, t, kicker=""):
    add_top_tag(slide, t)
    if kicker:
        add_text(slide, kicker.upper(), 0.72, 1.15, 3.0, 0.35, 10, t["muted"], True)
    add_text(slide, title, 0.72, 1.48, 11.6, 0.95, 30, t["fg"], True)


def add_decor(slide, t):
    p = panel(slide, 11.9, 0.55, 0.7, 0.7, t["accent"])
    p.rotation = 45
    panel(slide, 12.35, 6.65, 0.35, 0.35, t["panel"])


def build_presentation(data: dict, output_path: str | Path, theme="red", custom_colors=None):
    t = resolve_palette(theme, custom_colors)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    for idx, s in enumerate(data.get("slides", [])):
        slide = prs.slides.add_slide(blank)
        fill_bg(slide, t["bg"])
        add_decor(slide, t)
        typ = s.get("type", "content")
        title = s.get("title", "")
        subtitle = s.get("subtitle", "")
        content = s.get("content") or []

        if typ == "title":
            add_text(slide, "GENERATED WITH AI", 0.78, 0.72, 3.5, 0.35, 11, t["accent"], True)
            add_text(slide, title, 0.78, 2.0, 10.7, 1.75, 44, t["fg"], True)
            add_text(slide, subtitle or data.get("description", ""), 0.82, 4.0, 8.8, 0.8, 20, t["muted"])
            panel(slide, 0.82, 5.55, 4.4, 0.08, t["accent"], False)
            add_text(slide, f"{len(data.get('slides', []))} SLIDES", 0.82, 5.86, 2.0, 0.3, 10, t["muted"], True)

        elif typ == "comparison":
            slide_title(slide, title, t, "Comparison")
            pairs = [
                (0.75, s.get("left_title") or "A", s.get("left_text") or ""),
                (6.9, s.get("right_title") or "B", s.get("right_text") or ""),
            ]
            for x, head, txt in pairs:
                panel(slide, x, 2.65, 5.65, 3.55, t["panel"])
                add_text(slide, head, x + 0.35, 3.0, 4.8, 0.5, 18, t["accent"], True)
                add_text(slide, txt, x + 0.35, 3.7, 4.8, 1.8, 23, t["fg"], True)

        elif typ == "cards":
            slide_title(slide, title, t, "Key ideas")
            cards = (content + ["Idea", "Context", "Result"])[:3]
            for i, txt in enumerate(cards):
                x = 0.72 + i * 4.15
                panel(slide, x, 2.65, 3.75, 3.25, t["panel"])
                add_text(slide, f"0{i + 1}", x + 0.3, 2.95, 0.9, 0.5, 13, t["accent"], True)
                add_text(slide, txt, x + 0.3, 3.72, 3.0, 1.15, 25, t["fg"], True)

        elif typ == "stats":
            slide_title(slide, title, t, "Metric")
            add_text(slide, s.get("stat_value") or "100%", 0.78, 2.7, 11.2, 1.45, 64, t["accent"], True)
            add_text(slide, s.get("stat_label") or (content[0] if content else ""), 0.82, 4.45, 8.8, 1.0, 22, t["fg"])

        elif typ == "steps":
            slide_title(slide, title, t, "Process")
            steps = content[:5] or ["Start", "Build", "Review", "Export"]
            width = 11.7 / len(steps)
            for i, txt in enumerate(steps):
                x = 0.78 + i * width
                panel(slide, x, 3.05, min(width - 0.25, 2.65), 1.75, t["panel"])
                add_text(slide, str(i + 1).zfill(2), x + 0.25, 3.3, 0.5, 0.3, 10, t["accent"], True)
                add_text(slide, txt, x + 0.25, 3.83, min(width - 0.65, 2.1), 0.7, 17, t["fg"], True)

        elif typ == "quote":
            slide_title(slide, title, t, "Quote")
            add_text(slide, "“", 0.72, 2.55, 1.0, 1.0, 60, t["accent"], True)
            add_text(slide, content[0] if content else subtitle, 1.5, 2.85, 9.9, 2.3, 30, t["fg"], True)

        elif typ == "content_image":
            slide_title(slide, title, t, "Visual")
            add_bullets(slide, content, 0.78, 2.7, 5.25, 3.5, t, 20)
            img = fetch_pexels_image(s.get("image_query", ""))
            if img:
                try:
                    slide.shapes.add_picture(img, Inches(6.5), Inches(2.55), width=Inches(5.85), height=Inches(3.75))
                except Exception:
                    panel(slide, 6.5, 2.55, 5.85, 3.75, t["panel"])
            else:
                panel(slide, 6.5, 2.55, 5.85, 3.75, t["panel"])
                add_text(slide, s.get("image_query") or "IMAGE AREA", 6.9, 3.75, 5.0, 0.8, 16, t["muted"], True, PP_ALIGN.CENTER)

        elif typ == "ending":
            add_text(slide, "END", 0.8, 0.8, 1.0, 0.35, 11, t["accent"], True)
            add_text(slide, title or "Rahmat", 0.8, 2.2, 10.5, 1.2, 46, t["fg"], True)
            add_text(slide, subtitle, 0.82, 3.75, 9.0, 1.0, 21, t["muted"])
            panel(slide, 0.82, 5.5, 3.5, 0.08, t["accent"], False)

        else:
            slide_title(slide, title, t, "Overview")
            panel(slide, 0.72, 2.55, 11.6, 3.85, t["panel"])
            add_bullets(slide, content, 1.05, 2.95, 10.9, 3.0, t, 21)

        add_text(slide, f"{idx + 1:02d}", 12.18, 7.0, 0.45, 0.25, 9, t["muted"], True, PP_ALIGN.RIGHT)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return str(output_path)
